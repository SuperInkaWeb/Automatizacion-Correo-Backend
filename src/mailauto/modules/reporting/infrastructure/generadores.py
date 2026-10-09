"""
Generadores de reporte: Excel y CSV.

Proposito
    Escribir las filas en un fichero que un contador pueda abrir y
    usar directamente.

Dependencias
    openpyxl para Excel, biblioteca estandar para CSV.

Decisiones de diseño
    1. openpyxl en modo `write_only`. El modo normal mantiene todas
       las celdas en memoria hasta guardar; con treinta mil filas eso
       son cientos de megabytes en el worker. El modo de escritura
       vuelca fila a fila.

    2. El importe se escribe como NUMERO con formato de celda, no como
       texto. Un Excel donde la columna de importes son cadenas no
       suma, y lo primero que hace quien lo recibe es seleccionar la
       columna para ver el total.

    3. El CSV se escribe con BOM UTF-8. Sin el, Excel en Windows
       interpreta el fichero en la codificacion local y las tildes
       salen rotas: es el motivo numero uno de quejas con exportaciones
       en español.
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Final

from mailauto.modules.reporting.domain.ports import (
    COLUMNAS,
    FilaDeReporte,
    FormatoDeReporte,
    GeneradorDeReporte,
)

# Formato contable peruano: separador de miles y dos decimales.
_FORMATO_IMPORTE: Final = "#,##0.00"
# Columnas monetarias: se escriben como numero (no texto) para que la
# columna sume al seleccionarla.
_COLUMNAS_NUMERICAS: Final[frozenset[str]] = frozenset(
    {"monto_alquiler", "tributo_resultante", "importe_pagado", "intereses_moratorios"}
)
_ANCHOS: Final[dict[str, int]] = {
    "ruc_contribuyente": 14,
    "nombre_contribuyente": 38,
    "tipo_doc_inquilino": 16,
    "ruc_inquilino": 14,
    "nombre_inquilino": 38,
    "tipo_de_bien": 12,
    "periodo": 10,
    "monto_alquiler": 14,
    "tributo_resultante": 16,
    "importe_pagado": 14,
    "intereses_moratorios": 16,
    "moneda": 8,
    "fecha_de_pago": 13,
    "numero_de_operacion": 18,
    "estado": 18,
    "revision": 14,
    "archivo_origen": 34,
}


class GeneradorExcel(GeneradorDeReporte):
    @property
    def formato(self) -> FormatoDeReporte:
        return FormatoDeReporte.EXCEL

    def generar(self, filas: list[FilaDeReporte], titulo: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        # `write_only`: vuelca fila a fila en lugar de mantener la hoja
        # entera en memoria.
        libro = Workbook(write_only=True)
        hoja = libro.create_sheet(title=titulo[:31] or "Reporte")

        hoja.freeze_panes = "A2"
        for indice, (clave, _) in enumerate(COLUMNAS, start=1):
            hoja.column_dimensions[get_column_letter(indice)].width = _ANCHOS.get(clave, 18)

        negrita = Font(bold=True, color="FFFFFF")
        fondo = PatternFill("solid", fgColor="1F4E79")
        centrado = Alignment(horizontal="center", vertical="center")

        from openpyxl.cell import WriteOnlyCell

        cabecera = []
        for _, etiqueta in COLUMNAS:
            celda = WriteOnlyCell(hoja, value=etiqueta)
            celda.font = negrita
            celda.fill = fondo
            celda.alignment = centrado
            cabecera.append(celda)
        hoja.append(cabecera)

        for fila in filas:
            hoja.append(self._a_celdas(hoja, fila))

        bufer = io.BytesIO()
        libro.save(bufer)
        return bufer.getvalue()

    @staticmethod
    def _a_celdas(hoja: object, fila: FilaDeReporte) -> list[object]:
        from openpyxl.cell import WriteOnlyCell

        celdas: list[object] = []
        for clave, _ in COLUMNAS:
            valor = getattr(fila, clave)
            celda = WriteOnlyCell(hoja, value=valor)

            if clave in _COLUMNAS_NUMERICAS and valor:
                # Numero de verdad: asi la columna suma al seleccionarla.
                try:
                    celda.value = float(Decimal(valor))
                    celda.number_format = _FORMATO_IMPORTE
                except (InvalidOperation, ValueError):
                    # Importe ilegible: se deja el texto tal cual en
                    # lugar de perderlo o poner un cero enganoso.
                    celda.value = valor

            celdas.append(celda)
        return celdas


class GeneradorCsv(GeneradorDeReporte):
    @property
    def formato(self) -> FormatoDeReporte:
        return FormatoDeReporte.CSV

    def generar(self, filas: list[FilaDeReporte], titulo: str) -> bytes:
        bufer = io.StringIO(newline="")
        escritor = csv.writer(bufer, delimiter=";", quoting=csv.QUOTE_MINIMAL)

        escritor.writerow([etiqueta for _, etiqueta in COLUMNAS])
        for fila in filas:
            escritor.writerow([getattr(fila, clave) for clave, _ in COLUMNAS])

        # BOM UTF-8: sin el, Excel en Windows lee el fichero en la
        # codificacion local y las tildes salen rotas.
        return bufer.getvalue().encode("utf-8-sig")
